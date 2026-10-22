import base64
import importlib.util
import io
import re
import sys
import tempfile
import types
import unittest
import zipfile
from pathlib import Path

PACOTE = Path(__file__).resolve().parents[1]
SKILL = PACOTE.parent / 'gerar-minuta-impugnacao'
sys.path.insert(0, str(SKILL / 'scripts'))
sys.path.insert(0, str(PACOTE / 'build'))

from docx import Document

import gerar_minuta
import montar_tool

# O backend do Open WebUI fornece FastAPI. Este substituto deixa os testes de conversão rodarem também no ambiente
# isolado do pacote, sem alterar a integração real.
try:
    import fastapi.responses  # noqa: F401
except ModuleNotFoundError:
    fastapi = types.ModuleType('fastapi')
    responses = types.ModuleType('fastapi.responses')

    class HTMLResponse:
        def __init__(self, content, headers=None):
            self.body = content.encode('utf-8')
            self.headers = {key.lower(): value for key, value in (headers or {}).items()}

    responses.HTMLResponse = HTMLResponse
    fastapi.responses = responses
    sys.modules['fastapi'] = fastapi
    sys.modules['fastapi.responses'] = responses

TOOL = PACOTE / 'tools' / 'gerar_minuta_impugnacao.py'
spec = importlib.util.spec_from_file_location('gerar_minuta_impugnacao_tool', TOOL)
tool_module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(tool_module)

# Dados fictícios.
MARKDOWN = '''## IMPUGNAÇÃO AO EDITAL

**SERVFAZ SERVIÇOS LTDA.**, inscrita no CNPJ nº [Inserir CNPJ], com sede em [Inserir Endereço Completo],
representada por [Inserir Nome], RG [Inserir RG], CPF [Inserir CPF].

### I. DA TEMPESTIVIDADE

Conforme o item 13.1 do Edital [1], protocolada em [Data do Protocolo].

1. O **CONHECIMENTO** da impugnação;
2. No mérito, o **ACOLHIMENTO**:
   - **a)** A exclusão do item 1;
   - **b)** A adequação do item 2;
3. A republicação do edital;
'''
DADOS = {'CNPJ': '00.000.000/0001-00', 'ENDERECO': 'Rua Exemplo, n.º 1, Exemplópolis – UF'}


def docx_de(resposta):
    html = resposta[0].body.decode()
    dados = re.search(r'const b64 = "([^"]+)"', html)[1]
    return Document(io.BytesIO(base64.b64decode(dados))), html


class ToolTests(unittest.TestCase):
    def setUp(self):
        self.tool = tool_module.Tools()

    def preenchido(self):
        for chave, valor in DADOS.items():
            setattr(self.tool.valves, chave, valor)

    def test_returns_inline_html_with_download_and_context(self):
        self.preenchido()
        resposta = self.tool.gerar_minuta_impugnacao(MARKDOWN, data='2026-03-09', nome_arquivo='Impugnação UFPI 90009')
        page, contexto = resposta
        self.assertEqual(page.headers['content-disposition'], 'inline')
        self.assertTrue(contexto['ok'])
        self.assertEqual(contexto['arquivo'], 'Impugnacao-UFPI-90009.docx')
        self.assertEqual([campo['marcador'] for campo in contexto['campos_em_branco']],
                         ['[Inserir Nome]', '[Inserir RG]', '[Inserir CPF]'])
        doc, html = docx_de(resposta)
        self.assertIn('download="Impugnacao-UFPI-90009.docx"', html)
        textos = [p.text for p in doc.paragraphs]
        self.assertTrue(any('00.000.000/0001-00' in t and 'Rua Exemplo' in t for t in textos))
        self.assertTrue(any('[Inserir Nome]' in t and '[Inserir CPF]' in t for t in textos))
        self.assertTrue(any('9 de março de 2026' in t and '[1]' in t for t in textos))

    def test_missing_data_listed_and_highlighted(self):
        resposta = self.tool.gerar_minuta_impugnacao(MARKDOWN)  # sem valves nem data
        _, contexto = resposta
        marcadores = [c['marcador'] for c in contexto['campos_em_branco']]
        self.assertEqual(marcadores, ['[Inserir CNPJ]', '[Inserir Endereço Completo]', '[Inserir Nome]',
                                      '[Inserir RG]', '[Inserir CPF]', '[Data do Protocolo]'])
        doc, html = docx_de(resposta)
        self.assertIn('6 campo(s) em branco', html)
        marcados = [r.text for p in doc.paragraphs for r in p.runs if r.font.highlight_color is not None]
        self.assertEqual(marcados, marcadores)

    def test_legacy_tool_does_not_store_signer_personal_data_in_valves(self):
        self.assertFalse(hasattr(self.tool.valves, 'NOME'))
        self.assertFalse(hasattr(self.tool.valves, 'RG'))
        self.assertFalse(hasattr(self.tool.valves, 'CPF'))

    def test_errors_come_back_as_dict(self):
        for texto, data, trecho in (('Ver catálogo de fontes.', None, 'referência interna'),
                                    ('```\ncódigo\n```', None, 'não trata'), ('   ', None, 'vazio'),
                                    ('Texto.', '2026-13-01', 'data inválida')):
            resposta = self.tool.gerar_minuta_impugnacao(texto, data=data)
            self.assertIsInstance(resposta, dict)
            self.assertFalse(resposta['ok'])
            self.assertIn(trecho, resposta['erro'])

    def test_file_name_is_sanitized(self):
        seguro = tool_module.nome_seguro
        self.assertEqual(seguro(None), 'minuta-impugnacao.docx')
        self.assertEqual(seguro('../../etc/passwd'), 'etc-passwd.docx')
        self.assertEqual(seguro('a"b<c>.DOCX'), 'a-b-c.docx')

    def test_embedded_timbre_keeps_header_footer_and_watermark(self):
        with zipfile.ZipFile(io.BytesIO(base64.b64decode(tool_module._TIMBRE_B64))) as z:
            nomes = z.namelist()
            rodape = z.read('word/footer1.xml').decode()
        self.assertEqual(len([n for n in nomes if n.startswith('word/media/')]), 4)
        self.assertIn('<wp:anchor', rodape)
        self.assertIn('servfazpi', rodape)
        doc = Document(io.BytesIO(base64.b64decode(tool_module._TIMBRE_B64)))
        self.assertTrue(doc.sections[0].header.paragraphs)

    def test_same_output_as_the_skill_script(self):
        with tempfile.TemporaryDirectory() as tmp:
            fonte = Path(tmp) / 'peca.md'
            fonte.write_text(MARKDOWN, encoding='utf-8')
            original = gerar_minuta.convert(fonte, dados=Path(tmp) / 'sem-dados.json')
            esperado = Document(str(original.docx))
        doc, _ = docx_de(self.tool.gerar_minuta_impugnacao(MARKDOWN))
        resumo = lambda d: [(p.style.name, p.text, p.paragraph_format.left_indent) for p in d.paragraphs]
        self.assertEqual(resumo(doc), resumo(esperado))

    def test_tool_is_up_to_date_with_its_sources(self):
        self.assertEqual(TOOL.read_text(encoding='utf-8'), montar_tool.montar(),
                         'tools/gerar_minuta_impugnacao.py está defasada: rode build/montar_tool.py')

    def test_header_requirements_are_valid_after_openwebui_split(self):
        import ast
        from packaging.requirements import Requirement
        cabecalho = ast.get_docstring(ast.parse(TOOL.read_text(encoding='utf-8')))
        linha = next(l for l in cabecalho.splitlines() if l.startswith('requirements:'))
        pacotes = [Requirement(p.strip()) for p in linha.split(':', 1)[1].split(',') if p.strip()]
        self.assertEqual([p.name for p in pacotes], ['python-docx'])
        self.assertTrue(all(str(p.specifier).startswith('==') for p in pacotes))


if __name__ == '__main__':
    unittest.main()
