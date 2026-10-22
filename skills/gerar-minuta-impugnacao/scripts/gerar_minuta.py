#!/usr/bin/env python3
"""Converte uma minuta aprovada em Markdown no .docx timbrado da Servfaz.

Este módulo é a fonte canônica do conversor. A Action do Open WebUI e a Tool
legada incorporam somente o trecho entre os marcadores ``núcleo``. A conversão
é deliberadamente determinística: não redige, corrige, assina, protocola nem
envia a peça.

Uso:
  gerar_minuta.py ENTRADA.md [--saida MINUTA.docx] [--dados DADOS.json] [--data AAAA-MM-DD]
"""
import argparse
import json
import re
import sys
import unicodedata
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


# === núcleo compartilhado: este bloco é incorporado nas integrações Open WebUI ===
class Refusal(Exception):
    """Entrada que não pode virar minuta; a mensagem diz o que corrigir."""


# O catálogo de fontes é interno da triagem: nenhum ID, marcador [ID · status] ou menção a ele ou aos arquivos do
# estudo chega à minuta. "Catálogo" sozinho não conta: um edital pode citar o Catálogo de Serviços.
INTERNAL = re.compile(r'\b(?:EXT|INT|NOVA)-[A-Z0-9]|\[[^\]\n]*·[^\]\n]*\]|cat[aá]logo[ -]de[ -]fontes|'
                      r'(?:catalogo|indice)-fontes|references/estudo|\b[a-z]+/(?:fontes|analise|estrutura)\.md', re.I)

# Nomes em ASCII: o nome do estilo vira o styleId no XML, e leitores de .docx tropeçam em acento ali.
STYLES = {'corpo': 'Peca Corpo', 'titulo': 'Peca Titulo', 'subtitulo': 'Peca Subtitulo', 'citacao': 'Peca Citacao',
          'alinea': 'Peca Alinea', 'centro': 'Peca Centro', 'tabela': 'Peca Tabela'}
ALINEA_NUMBERING = 'PecaAlinea'
MESES = ('janeiro', 'fevereiro', 'março', 'abril', 'maio', 'junho', 'julho', 'agosto', 'setembro', 'outubro',
         'novembro', 'dezembro')
# Rótulo do "[Inserir ...]" (sem acento, minúsculo) -> chave de dados-servfaz.json.
CAMPOS = {'cnpj': 'cnpj', 'endereco': 'endereco', 'endereco completo': 'endereco', 'cpf': 'cpf', 'rg': 'rg',
          'nome': 'nome', 'nome do representante': 'nome', 'representante': 'nome'}
CAMPO = re.compile(r'\[inserir ([^\]\n]+)\]', re.I)
CAMPO_DATA = re.compile(r'\[data(?: do protocolo| da pe[cç]a)?\]', re.I)
ITEM = re.compile(r'^\d+\.\s+(.*)$')
BOLD_LINE = re.compile(r'^\*\*[^*]+\*\*$')
INLINE = re.compile(r'(\*\*.+?\*\*|(?<![\w*])\*(?!\*)[^*]+?\*(?!\*)|(?<![\w*])_(?!_).+?(?<!_)_(?!\w))')
HEADING = re.compile(r'^(#{1,6})\s+(.*)$')
ROMAN = re.compile(r'^[IVXL]+\.\s+')
BULLET = re.compile(r'^\s*[-*+]\s+(.*)$')
RULE = re.compile(r'^(?:-{3,}|\*{3,})$')
MARKER = re.compile(r'^\**[a-z0-9]+\)\**\s')
SUB_INDENT = Cm(2.25)
FRONT = re.compile(r'\A---\n.*?\n---\n', re.S)
CENTER_ON, CENTER_OFF = '<!-- centralizado -->', '<!-- /centralizado -->'

# Campo em branco: não barra a geração; sai em destaque no .docx e entra na lista de pendências.
PENDENTE = re.compile(r'\{\{[^}\n]*\}\}|\[(?:inserir|preencher|informar|data|nome|cnpj|cpf|endere[cç]o)\b[^\]\n]*\]', re.I)
# Marcação que o conversor não trata. Remissões como [1] passam: vão ao .docx como estão.
UNSUPPORTED = re.compile(r'^\s*```|^#{5,}\s|!\[[^\]\n]*\]\(|\[[^\]\n]+\]\([^)\n]+\)|</?[a-z][^>\n]*>|~~', re.I)


def add_text(paragraph, text, **fmt):
    """Texto em um ou mais trechos; o campo em branco sai com destaque amarelo."""
    for piece in re.split(f'({PENDENTE.pattern})', text, flags=re.I):
        if not piece:
            continue
        run = paragraph.add_run(piece)
        for name, value in fmt.items():
            setattr(run, name, value)
        if PENDENTE.fullmatch(piece):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def add_runs(paragraph, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            add_text(paragraph, part[2:-2], bold=True)
        elif part[0] in '*_' and part[-1] == part[0] and len(part) > 2:
            add_text(paragraph, part[1:-1], italic=True)
        else:
            add_text(paragraph, part)
    return paragraph


def new_list(doc):
    """Nova numeração a), b) que recomeça em a): um w:num por lista, com startOverride."""
    numbering = doc.part.numbering_part.element
    abstract_id = None
    for abstract in numbering.findall(qn('w:abstractNum')):
        name = abstract.find(qn('w:name'))
        if name is not None and name.get(qn('w:val')) == ALINEA_NUMBERING:
            abstract_id = abstract.get(qn('w:abstractNumId'))
    if abstract_id is None:
        raise Refusal(f'o timbre não tem a numeração {ALINEA_NUMBERING}: gere-o com criar_timbre.py.')
    num_id = max([int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))], default=0) + 1
    num = OxmlElement('w:num')
    num.set(qn('w:numId'), str(num_id))
    ref = OxmlElement('w:abstractNumId')
    ref.set(qn('w:val'), abstract_id)
    override = OxmlElement('w:lvlOverride')
    override.set(qn('w:ilvl'), '0')
    start = OxmlElement('w:startOverride')
    start.set(qn('w:val'), '1')
    override.append(start)
    num.extend([ref, override])
    numbering.append(num)
    return num_id


def number_paragraph(paragraph, num_id):
    numpr = OxmlElement('w:numPr')
    for tag, value in (('w:ilvl', '0'), ('w:numId', str(num_id))):
        node = OxmlElement(tag)
        node.set(qn('w:val'), value)
        numpr.append(node)
    paragraph._p.get_or_add_pPr().append(numpr)


def add_table(doc, lines):
    rows = [[cell.strip() for cell in line.strip().strip('|').split('|')] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r':?-+:?', cell) for cell in row)]
    if not rows or not rows[0]:
        raise Refusal('a tabela Markdown está vazia.')
    width = len(rows[0])
    if any(len(row) != width for row in rows):
        raise Refusal('a tabela Markdown tem linhas com quantidades diferentes de colunas.')
    table = doc.add_table(rows=len(rows), cols=width)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge = OxmlElement(f'w:{side}')
        for key, value in (('val', 'single'), ('sz', '4'), ('space', '0'), ('color', '000000')):
            edge.set(qn(f'w:{key}'), value)
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    for row_number, row in enumerate(rows):
        for column, text in enumerate(row):
            paragraph = table.cell(row_number, column).paragraphs[0]
            paragraph.style = STYLES['tabela']
            add_runs(paragraph, f'**{text}**' if row_number == 0 and text else text)
    header = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(header)
    doc.add_paragraph(style=STYLES['corpo'])


def add_sub_items(doc, subs):
    for sub in subs:
        text = sub if MARKER.match(sub) else f'– {sub}'
        paragraph = add_runs(doc.add_paragraph(style=STYLES['alinea']), text)
        paragraph.paragraph_format.left_indent = SUB_INDENT


def close_last(items):
    """O último item do pedido termina em ponto final, mesmo que o texto traga ';'."""
    last = items[-1]
    if last[1]:
        last[1][-1] = re.sub(r';\s*$', '.', last[1][-1])
    else:
        last[0] = re.sub(r';\s*$', '.', last[0])


def build(doc, text):
    lines = text.split('\n')
    centered, index = False, 0
    blank = ('', CENTER_ON, CENTER_OFF)
    while index < len(lines):
        line = lines[index].strip()
        following = lines[index + 1].strip() if index + 1 < len(lines) else ''
        index += 1
        if not line or RULE.match(line):
            continue
        if line in (CENTER_ON, CENTER_OFF):
            centered = line == CENTER_ON
            continue
        heading = HEADING.match(line)
        if heading:
            level, title = len(heading[1]), heading[2].strip()
            if level == 2:
                paragraph = add_runs(doc.add_paragraph(style=STYLES['centro']), title)
                for run in paragraph.runs:
                    run.bold = True
            elif level == 4:
                add_runs(doc.add_paragraph(style=STYLES['subtitulo']), title)
            else:
                add_runs(doc.add_paragraph(style=STYLES['titulo']), ROMAN.sub('', title))
            continue
        if line.startswith('|'):
            block = [line]
            while index < len(lines) and lines[index].strip().startswith('|'):
                block.append(lines[index].strip())
                index += 1
            add_table(doc, block)
            continue
        if ITEM.match(line):
            items = [[ITEM.match(line)[1], []]]
            while index < len(lines):
                candidate = lines[index]
                if ITEM.match(candidate.strip()):
                    items.append([ITEM.match(candidate.strip())[1], []])
                elif BULLET.match(candidate):
                    items[-1][1].append(BULLET.match(candidate)[1].strip())
                else:
                    break
                index += 1
            close_last(items)
            num_id = new_list(doc)
            for item, subs in items:
                number_paragraph(add_runs(doc.add_paragraph(style=STYLES['alinea']), item), num_id)
                add_sub_items(doc, subs)
            continue
        if BULLET.match(lines[index - 1]):
            bullets = [BULLET.match(lines[index - 1])[1].strip()]
            while index < len(lines) and BULLET.match(lines[index]):
                bullets.append(BULLET.match(lines[index])[1].strip())
                index += 1
            if all(MARKER.match(bullet) for bullet in bullets):
                add_sub_items(doc, bullets)
            else:
                items = [[bullet, []] for bullet in bullets]
                close_last(items)
                num_id = new_list(doc)
                for item, _ in items:
                    number_paragraph(add_runs(doc.add_paragraph(style=STYLES['alinea']), item), num_id)
            continue
        if line.startswith('>'):
            add_runs(doc.add_paragraph(style=STYLES['citacao']), line[1:].strip())
            continue
        style = STYLES['centro'] if centered else STYLES['subtitulo'] if BOLD_LINE.match(line) else STYLES['corpo']
        paragraph = add_runs(doc.add_paragraph(style=style), line)
        if following and following not in blank and not RULE.match(following):
            paragraph.paragraph_format.space_after = Pt(0)


def preencher(text, dados, when=None):
    """Preenche só os campos que os dados conhecem; os demais permanecem destacados para revisão humana."""
    def campo(match):
        label = unicodedata.normalize('NFKD', match[1]).encode('ascii', 'ignore').decode().lower().strip()
        return dados.get(CAMPOS.get(label)) or match[0]

    text = CAMPO.sub(campo, text)
    if when:
        text = CAMPO_DATA.sub(f'{when.day} de {MESES[when.month - 1]} de {when.year}', text)
    return text


def check_clean(text, label):
    """A trava: a minuta não sai com referência interna da triagem nem marcação que o conversor não trata."""
    problems = []
    for number, line in enumerate(text.split('\n'), 1):
        for regex, what in ((INTERNAL, 'referência interna'), (UNSUPPORTED, 'marcação que o conversor não trata')):
            found = regex.search(line)
            if found:
                problems.append(f'linha {number}: {what} "{found[0].strip()}"')
    if problems:
        extra = f'; e mais {len(problems) - 10}' if len(problems) > 10 else ''
        raise Refusal(f'{label} não pode sair assim: ' + '; '.join(problems[:10]) + extra +
                      '. Corrija o Markdown e gere de novo.')
    return text


def pendencias(text):
    """Campos em branco que sobraram, como (linha, marcador), para o analista preencher no Word."""
    return [(number, found[0]) for number, line in enumerate(text.split('\n'), 1) for found in PENDENTE.finditer(line)]


def render(text, timbre, dados=None, when=None, label='a minuta'):
    """Markdown -> (Document, pendências). ``timbre`` aceita caminho ou arquivo binário aberto."""
    text = preencher(unicodedata.normalize('NFC', text), dados or {}, when)
    check_clean(text, label)
    text = FRONT.sub('', text, count=1)
    doc = Document(timbre)
    build(doc, text)
    return doc, pendencias(text)
# === fim do núcleo ===


TIMBRE = Path(__file__).resolve().parents[1] / 'assets' / 'timbre-servfaz.docx'


@dataclass(frozen=True)
class Conversion:
    docx: Path
    campos_em_branco: list[tuple[int, str]]


def load_dados(source):
    if not source:
        return {}
    if isinstance(source, dict):
        return source
    path = Path(source)
    if not path.is_file():
        return {}
    try:
        data = json.loads(path.read_text(encoding='utf-8'))
    except json.JSONDecodeError as exc:
        raise Refusal(f'dados inválidos em {path.name}: JSON inválido.') from exc
    if not isinstance(data, dict):
        raise Refusal(f'dados inválidos em {path.name}: esperado um objeto JSON.')
    return {str(key).lower(): str(value) for key, value in data.items() if value is not None}


def convert(source, target=None, timbre=TIMBRE, dados=None, when=None):
    source = Path(source)
    if not source.is_file():
        raise Refusal(f'arquivo Markdown não encontrado: {source}.')
    if not Path(timbre).is_file():
        raise Refusal(f'timbre não encontrado ({timbre}): gere-o com criar_timbre.py.')
    if isinstance(when, str):
        try:
            when = date.fromisoformat(when)
        except ValueError as exc:
            raise Refusal('data inválida: use AAAA-MM-DD.') from exc
    doc, missing = render(source.read_text(encoding='utf-8'), str(timbre), load_dados(dados), when, source.name)
    output = Path(target) if target else source.with_suffix('.docx')
    doc.save(str(output))
    return Conversion(docx=output, campos_em_branco=missing)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('entrada', type=Path, help='minuta aprovada em Markdown')
    parser.add_argument('--saida', type=Path, help='destino (padrão: mesmo nome, .docx)')
    parser.add_argument('--dados', type=Path, help='JSON opcional com cnpj, endereco, nome, rg e cpf')
    parser.add_argument('--data', help='data da peça em AAAA-MM-DD; ausente, permanece em destaque amarelo')
    args = parser.parse_args(argv)
    try:
        result = convert(args.entrada, args.saida, dados=args.dados, when=args.data)
    except (OSError, Refusal) as exc:
        print(f'erro: {exc}', file=sys.stderr)
        return 1
    print(result.docx)
    if result.campos_em_branco:
        print('campos pendentes: ' + ', '.join(f'linha {line}: {field}' for line, field in result.campos_em_branco),
              file=sys.stderr)
    return 0


if __name__ == '__main__':
    sys.exit(main())
